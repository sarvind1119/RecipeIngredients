"""Application-level tests: auth, recipe entry, scaling, snapshots, exports.

Each test gets a fresh throwaway database. Run from the project root:
    python -m pytest tests/ -v
"""

import io
import os
import sqlite3
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import db  # noqa: E402

BIRYANI = {
    "name": "Chicken Biryani",
    "category": "Main Course",
    "base_persons": "10",
    "notes": "Dum style.",
    "ingredient_name": ["Chicken", "Basmati Rice", "Onion", "Ghee", "Salt", "Eggs"],
    "ingredient_quantity": ["3", "2", "1.5", "500", "100", "5"],
    "ingredient_unit": ["kg", "kg", "kg", "g", "g", "nos"],
}


@pytest.fixture()
def application(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", str(tmp_path / "mess.db"))
    import app as appmod

    appmod.app.config["TESTING"] = True
    db.init_db()
    return appmod.app


@pytest.fixture()
def admin(application):
    """A logged-in admin that has already cleared the forced password change."""
    client = application.test_client()
    client.post("/login", data={"username": "admin", "password": "admin123"})
    client.post(
        "/change-password",
        data={"current_password": "admin123", "new_password": "mess2026",
              "confirm_password": "mess2026"},
    )
    return client


@pytest.fixture()
def staff(application, admin):
    admin.post("/users", data={"username": "cook", "full_name": "Head Cook",
                               "password": "cook2026", "role": "staff"})
    client = application.test_client()
    client.post("/login", data={"username": "cook", "password": "cook2026"})
    client.post(
        "/change-password",
        data={"current_password": "cook2026", "new_password": "cook12345",
              "confirm_password": "cook12345"},
    )
    return client


def _requisition_items(requisition_id=1):
    conn = db.connect()
    rows = {
        r["ingredient_name"]: r
        for r in conn.execute(
            "SELECT * FROM requisition_items WHERE requisition_id = ?", (requisition_id,)
        )
    }
    conn.close()
    return rows


class TestFirstRun:
    def test_seeds_a_bootstrap_admin(self, application):
        conn = db.connect()
        user = conn.execute("SELECT * FROM users").fetchone()
        conn.close()
        assert user["username"] == "admin"
        assert user["must_change_password"] == 1

    def test_seed_is_not_repeated(self, application):
        assert db.init_db() is False

    def test_default_password_forces_a_change(self, application):
        client = application.test_client()
        r = client.post("/login", data={"username": "admin", "password": "admin123"})
        assert "/change-password" in r.headers["Location"]

    def test_wrong_password_rejected(self, application):
        client = application.test_client()
        r = client.post("/login", data={"username": "admin", "password": "wrong"},
                        follow_redirects=True)
        assert b"Invalid username or password" in r.data


class TestPasswordGate:
    """Regression: the forced change must gate admin routes, not only the
    login_required ones. A seeded admin could otherwise reach /dishes/new by
    typing the URL while still on the default password."""

    @pytest.mark.parametrize("path", ["/dishes/new", "/users", "/admin/backup", "/dishes"])
    def test_routes_redirect_until_password_changed(self, application, path):
        client = application.test_client()
        client.post("/login", data={"username": "admin", "password": "admin123"})
        r = client.get(path)
        assert r.status_code == 302
        assert "/change-password" in r.headers["Location"]

    def test_recipe_cannot_be_created_before_change(self, application):
        client = application.test_client()
        client.post("/login", data={"username": "admin", "password": "admin123"})
        client.post("/dishes/new", data=BIRYANI)

        conn = db.connect()
        count = conn.execute("SELECT COUNT(*) FROM dishes").fetchone()[0]
        conn.close()
        assert count == 0


class TestRoles:
    def test_staff_can_create_a_recipe(self, staff):
        r = staff.post("/dishes/new", data=BIRYANI, follow_redirects=True)
        assert r.status_code == 200
        assert b"Chicken Biryani" in r.data

    def test_staff_can_edit_and_delete(self, admin, staff):
        admin.post("/dishes/new", data=BIRYANI)
        amended = dict(BIRYANI)
        amended["name"] = "Chicken Biryani (staff edit)"
        r = staff.post("/dishes/1/edit", data=amended, follow_redirects=True)
        assert r.status_code == 200
        assert b"staff edit" in r.data
        assert staff.post("/dishes/1/delete", follow_redirects=True).status_code == 200

    def test_staff_cannot_manage_users(self, staff):
        assert staff.get("/users").status_code == 403

    def test_staff_cannot_download_the_backup(self, staff):
        assert staff.get("/admin/backup").status_code == 403

    def test_staff_can_still_calculate_and_print(self, admin, staff):
        admin.post("/dishes/new", data=BIRYANI)
        r = staff.post("/calculate", data={"dish_id": "1", "persons_required": "90",
                                           "meal_type": "Lunch", "meal_date": "2026-08-10"},
                       follow_redirects=True)
        assert r.status_code == 200
        assert staff.get("/requisition/1/pdf").status_code == 200


class TestRecipeEntry:
    def test_blank_rows_are_dropped_silently(self, admin):
        payload = dict(BIRYANI)
        payload["ingredient_name"] = BIRYANI["ingredient_name"] + ["", ""]
        payload["ingredient_quantity"] = BIRYANI["ingredient_quantity"] + ["", ""]
        payload["ingredient_unit"] = BIRYANI["ingredient_unit"] + ["kg", "kg"]

        r = admin.post("/dishes/new", data=payload, follow_redirects=True)
        assert b"saved" in r.data

        conn = db.connect()
        count = conn.execute("SELECT COUNT(*) FROM dish_ingredients").fetchone()[0]
        conn.close()
        assert count == 6

    def test_partially_filled_row_is_reported(self, admin):
        r = admin.post("/dishes/new", data={
            "name": "Partial", "category": "Snacks", "base_persons": "10", "notes": "",
            "ingredient_name": ["Flour", "Sugar"], "ingredient_quantity": ["1", ""],
            "ingredient_unit": ["kg", "kg"]}, follow_redirects=True)
        assert b"has no quantity" in r.data

    def test_duplicate_ingredient_rejected(self, admin):
        r = admin.post("/dishes/new", data={
            "name": "Dupe", "category": "Snacks", "base_persons": "10", "notes": "",
            "ingredient_name": ["Onion", "onion"], "ingredient_quantity": ["1", "2"],
            "ingredient_unit": ["kg", "kg"]}, follow_redirects=True)
        assert b"more than once" in r.data

    def test_duplicate_dish_name_rejected(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        r = admin.post("/dishes/new", data=BIRYANI, follow_redirects=True)
        assert b"already exists" in r.data

    def test_user_input_survives_a_validation_failure(self, admin):
        r = admin.post("/dishes/new", data={
            "name": "Half Typed Dish", "category": "Snacks", "base_persons": "",
            "notes": "", "ingredient_name": ["Flour"], "ingredient_quantity": ["2"],
            "ingredient_unit": ["kg"]}, follow_redirects=True)
        # Nobody should have to key a whole recipe in twice over one bad field.
        assert b"Half Typed Dish" in r.data
        assert b"Flour" in r.data

    def test_ingredient_names_are_case_folded_into_one_master_row(self, admin):
        admin.post("/dishes/new", data={
            "name": "Dish A", "category": "Snacks", "base_persons": "10", "notes": "",
            "ingredient_name": ["Onion"], "ingredient_quantity": ["1"],
            "ingredient_unit": ["kg"]})
        admin.post("/dishes/new", data={
            "name": "Dish B", "category": "Snacks", "base_persons": "10", "notes": "",
            "ingredient_name": ["onion"], "ingredient_quantity": ["2"],
            "ingredient_unit": ["kg"]})

        conn = db.connect()
        count = conn.execute("SELECT COUNT(*) FROM ingredients").fetchone()[0]
        conn.close()
        assert count == 1  # what makes cross-dish consolidation possible later

    def test_autocomplete_endpoint_returns_the_master_list(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        payload = admin.get("/api/ingredients?q=chick").get_json()
        assert any(item["name"] == "Chicken" for item in payload)


class TestScalingThroughTheApp:
    @pytest.fixture(autouse=True)
    def _recipe(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        self.admin = admin

    def test_the_requirement_document_example(self):
        self.admin.post("/calculate", data={"dish_id": "1", "persons_required": "90",
                                            "course_name": "FC 2026", "meal_type": "Lunch",
                                            "meal_date": "2026-08-10"})
        items = _requisition_items()
        assert items["Chicken"]["display_quantity"] == 27.0
        assert items["Chicken"]["display_unit"] == "kg"
        assert repr(items["Chicken"]["exact_quantity"]) == "27.0"

    def test_unit_promotion_end_to_end(self):
        self.admin.post("/calculate", data={"dish_id": "1", "persons_required": "90",
                                            "meal_type": "Lunch", "meal_date": "2026-08-10"})
        ghee = _requisition_items()["Ghee"]
        assert ghee["exact_quantity"] == 4500.0
        assert ghee["base_unit"] == "g"
        assert (ghee["display_quantity"], ghee["display_unit"]) == (4.5, "kg")

    def test_countable_rounds_up(self):
        self.admin.post("/calculate", data={"dish_id": "1", "persons_required": "65",
                                            "meal_type": "Dinner", "meal_date": "2026-08-11"})
        eggs = _requisition_items()["Eggs"]
        assert eggs["exact_quantity"] == 32.5
        assert eggs["display_quantity"] == 33.0

    def test_meal_type_is_required(self):
        r = self.admin.post("/calculate", data={"dish_id": "1", "persons_required": "90",
                                                "meal_type": "", "meal_date": "2026-08-10"},
                            follow_redirects=True)
        assert b"select a meal type" in r.data

    def test_zero_persons_rejected(self):
        r = self.admin.post("/calculate", data={"dish_id": "1", "persons_required": "0",
                                                "meal_type": "Lunch", "meal_date": "2026-08-10"},
                            follow_redirects=True)
        assert b"at least 1" in r.data


class TestSnapshotIntegrity:
    """The reason requisition_items exists at all."""

    def test_editing_a_recipe_does_not_alter_a_past_requisition(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        admin.post("/calculate", data={"dish_id": "1", "persons_required": "90",
                                       "meal_type": "Lunch", "meal_date": "2026-08-10"})

        amended = dict(BIRYANI)
        amended["ingredient_quantity"] = ["3.5", "2", "1.5", "500", "100", "5"]
        admin.post("/dishes/1/edit", data=amended)

        assert _requisition_items()["Chicken"]["display_quantity"] == 27.0

    def test_renaming_a_dish_does_not_alter_a_past_requisition(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        admin.post("/calculate", data={"dish_id": "1", "persons_required": "90",
                                       "meal_type": "Lunch", "meal_date": "2026-08-10"})

        renamed = dict(BIRYANI, name="Mutton Biryani")
        admin.post("/dishes/1/edit", data=renamed)

        conn = db.connect()
        snapshot = conn.execute(
            "SELECT dish_name_snapshot FROM requisitions WHERE id = 1").fetchone()
        conn.close()
        assert snapshot["dish_name_snapshot"] == "Chicken Biryani"

    def test_soft_deleted_dish_keeps_its_requisitions_readable(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        admin.post("/calculate", data={"dish_id": "1", "persons_required": "90",
                                       "meal_type": "Lunch", "meal_date": "2026-08-10"})
        admin.post("/dishes/1/delete")

        assert admin.get("/requisition/1").status_code == 200
        assert admin.get("/requisition/1/pdf").status_code == 200
        assert b"Chicken Biryani" not in admin.get("/dishes").data


class TestSoftDeleteLifecycle:
    def test_soft_deleted_name_can_be_reclaimed_by_readding(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        admin.post("/dishes/1/delete")

        # Active-list clash no longer blocks; UNIQUE name is reclaimed on the
        # same dish_id so past requisitions stay linked to one recipe row.
        r = admin.post("/dishes/new", data=BIRYANI, follow_redirects=True)
        assert b"saved" in r.data

        conn = db.connect()
        rows = conn.execute(
            "SELECT id, is_active FROM dishes WHERE name = 'Chicken Biryani'"
        ).fetchall()
        conn.close()
        assert len(rows) == 1
        assert rows[0]["id"] == 1
        assert rows[0]["is_active"] == 1

    def test_restore_endpoint_reactivates_a_removed_dish(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        admin.post("/dishes/1/delete")
        # Flash text mentions the name; the active list table must not.
        list_html = admin.get("/dishes").data.decode()
        assert 'href="/dishes/1"' not in list_html
        assert "No recipes recorded yet." in list_html

        r = admin.post("/dishes/1/restore", follow_redirects=True)
        assert b"restored" in r.data
        list_html = admin.get("/dishes").data.decode()
        assert 'href="/dishes/1"' in list_html
        assert "Chicken Biryani" in list_html

    def test_calculate_rejects_soft_deleted_dish(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        admin.post("/dishes/1/delete")

        r = admin.post(
            "/calculate",
            data={
                "dish_id": "1",
                "persons_required": "90",
                "meal_type": "Lunch",
                "meal_date": "2026-08-10",
            },
            follow_redirects=True,
        )
        assert b"select a dish" in r.data.lower() or b"Please select a dish" in r.data

        conn = db.connect()
        count = conn.execute("SELECT COUNT(*) FROM requisitions").fetchone()[0]
        conn.close()
        assert count == 0


class TestTransactionalWrites:
    def test_failed_ingredient_insert_does_not_leave_an_empty_dish(self, admin, monkeypatch):
        """Recipe replace must roll back if a later write fails mid-loop."""
        admin.post("/dishes/new", data=BIRYANI)

        real_execute = db.execute
        calls = {"n": 0}

        def flaky_execute(sql, args=()):
            # After the dish UPDATE and DELETE dish_ingredients, fail on the
            # first dish_ingredients INSERT so a non-transactional path would
            # leave the recipe empty.
            if "INSERT INTO dish_ingredients" in sql:
                calls["n"] += 1
                if calls["n"] == 1:
                    raise sqlite3.OperationalError("simulated mid-save failure")
            return real_execute(sql, args)

        monkeypatch.setattr(db, "execute", flaky_execute)

        amended = dict(BIRYANI)
        amended["ingredient_quantity"] = ["3.5", "2", "1.5", "500", "100", "5"]
        with pytest.raises(sqlite3.OperationalError):
            admin.post("/dishes/1/edit", data=amended)

        conn = db.connect()
        count = conn.execute(
            "SELECT COUNT(*) FROM dish_ingredients WHERE dish_id = 1"
        ).fetchone()[0]
        chicken = conn.execute(
            """SELECT di.quantity FROM dish_ingredients di
                 JOIN ingredients i ON i.id = di.ingredient_id
                WHERE di.dish_id = 1 AND i.name = 'Chicken'"""
        ).fetchone()
        conn.close()
        # Original six rows still present; chicken still 3 kg, not half-edited.
        assert count == 6
        assert chicken["quantity"] == 3.0

    def test_failed_requisition_item_insert_leaves_no_partial_requisition(
        self, admin, monkeypatch
    ):
        admin.post("/dishes/new", data=BIRYANI)

        real_execute = db.execute

        def flaky_execute(sql, args=()):
            if "INSERT INTO requisition_items" in sql:
                raise sqlite3.OperationalError("simulated mid-requisition failure")
            return real_execute(sql, args)

        monkeypatch.setattr(db, "execute", flaky_execute)

        with pytest.raises(sqlite3.OperationalError):
            admin.post(
                "/calculate",
                data={
                    "dish_id": "1",
                    "persons_required": "90",
                    "meal_type": "Lunch",
                    "meal_date": "2026-08-10",
                },
            )

        conn = db.connect()
        assert conn.execute("SELECT COUNT(*) FROM requisitions").fetchone()[0] == 0
        assert conn.execute("SELECT COUNT(*) FROM requisition_items").fetchone()[0] == 0
        conn.close()


class TestIngredientRace:
    def test_get_or_create_recovers_from_unique_race(self, application, monkeypatch):
        """When the first SELECT misses a concurrent insert, IntegrityError re-resolves."""
        with application.app_context():
            first = db.get_or_create_ingredient("Cardamom", "g")
            db.commit()

            real_lookup = db._ingredient_id_by_name
            state = {"missed": False}

            def miss_once(conn, name):
                if not state["missed"] and name.casefold() == "cardamom":
                    state["missed"] = True
                    return None
                return real_lookup(conn, name)

            monkeypatch.setattr(db, "_ingredient_id_by_name", miss_once)
            again = db.get_or_create_ingredient("Cardamom", "kg")
            assert again == first
            assert state["missed"] is True


class TestExports:
    @pytest.fixture(autouse=True)
    def _requisition(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        admin.post("/calculate", data={"dish_id": "1", "persons_required": "90",
                                       "course_name": "FC 2026", "meal_type": "Lunch",
                                       "meal_date": "2026-08-10"})
        self.admin = admin

    def test_pdf_is_a_pdf(self):
        r = self.admin.get("/requisition/1/pdf")
        assert r.status_code == 200
        assert r.data[:4] == b"%PDF"

    def test_pdf_spans_pages_for_a_long_ingredient_list(self, admin):
        admin.post("/dishes/new", data={
            "name": "Grand Thali", "category": "Main Course", "base_persons": "10", "notes": "",
            "ingredient_name": [f"Item {i}" for i in range(1, 31)],
            "ingredient_quantity": ["1.25"] * 30, "ingredient_unit": ["kg"] * 30})
        admin.post("/calculate", data={"dish_id": "2", "persons_required": "120",
                                       "meal_type": "Special Event", "meal_date": "2026-09-01"})
        data = admin.get("/requisition/2/pdf").data
        assert data[:4] == b"%PDF"
        assert data.count(b"/Type /Page\n") >= 2 or b"Page 2" in data

    def test_excel_keeps_quantities_numeric_with_their_own_units(self):
        from openpyxl import load_workbook

        r = self.admin.get("/requisition/1/excel")
        ws = load_workbook(io.BytesIO(r.data)).active

        header = next(rw for rw in range(1, ws.max_row + 1)
                      if ws.cell(row=rw, column=1).value == "S.No")
        headings = [ws.cell(row=header, column=c).value for c in range(1, ws.max_column + 1)]
        assert "Exact Unit" in headings

        ghee = next([ws.cell(row=rw, column=c).value for c in range(1, 8)]
                    for rw in range(header + 1, ws.max_row + 1)
                    if ws.cell(row=rw, column=2).value == "Ghee")
        # Exact 4500 g and Required 4.5 kg sit in different units, so the exact
        # figure must carry its own - otherwise the row reads as 4500 kg.
        assert ghee[3] == 4500 and ghee[4] == "g"
        assert ghee[5] == 4.5 and ghee[6] == "kg"

    def test_screen_shows_the_exact_figure_with_its_unit(self):
        html = self.admin.get("/requisition/1").data.decode()
        assert "4500 g" in html


class TestBackup:
    def test_backup_includes_writes_made_moments_earlier(self, admin, tmp_path):
        """Under WAL a naive send_file('mess.db') can miss recent commits."""
        admin.post("/dishes/new", data=BIRYANI)
        r = admin.get("/admin/backup")
        assert r.status_code == 200

        path = tmp_path / "backup.db"
        path.write_bytes(r.data)
        conn = sqlite3.connect(path)
        names = [row[0] for row in conn.execute("SELECT name FROM dishes")]
        conn.close()
        assert "Chicken Biryani" in names


    def test_backup_leaves_no_copy_in_the_temp_folder(self, admin):
        """The snapshot holds every password hash - it must not be left behind."""
        import glob
        import tempfile

        pattern = os.path.join(tempfile.gettempdir(), "mess_backup_*")
        before = set(glob.glob(pattern))
        assert admin.get("/admin/backup").status_code == 200
        assert set(glob.glob(pattern)) == before


class TestUploadLimit:
    def test_max_content_length_is_set(self, application):
        """Enforced while the body arrives, so an oversized file is never
        written to the Mess PC's disk and only measured afterwards."""
        import app as appmod

        assert application.config["MAX_CONTENT_LENGTH"] == appmod.MAX_IMPORT_BYTES

    def test_oversized_upload_gets_the_413_page_not_a_traceback(self, admin, application):
        import app as appmod

        oversized = b"x" * (appmod.MAX_IMPORT_BYTES + 1024)
        r = admin.post(
            "/dishes/import",
            data={"document": (io.BytesIO(oversized), "huge.png")},
            content_type="multipart/form-data",
        )
        assert r.status_code == 413
        assert b"larger than 25 MB" in r.data


class TestDatabasePragmas:
    def test_wal_and_foreign_keys_are_enabled(self, application):
        conn = db.connect()
        assert conn.execute("PRAGMA journal_mode").fetchone()[0].lower() == "wal"
        assert conn.execute("PRAGMA foreign_keys").fetchone()[0] == 1
        conn.close()

    def test_foreign_keys_are_actually_enforced(self, application):
        conn = db.connect()
        with pytest.raises(sqlite3.IntegrityError):
            conn.execute("""INSERT INTO dish_ingredients (dish_id, ingredient_id, quantity, unit)
                            VALUES (9999, 8888, 1, 'kg')""")
            conn.commit()
        conn.close()

    def test_concurrent_writers_do_not_lock_each_other_out(self, application):
        import threading

        errors = []

        def writer(tag):
            try:
                conn = db.connect()
                for i in range(25):
                    conn.execute(
                        """INSERT INTO requisitions
                               (dish_name_snapshot, base_persons_snapshot, persons_required,
                                meal_type, meal_date)
                           VALUES (?, 10, 90, 'Lunch', '2026-08-10')""",
                        (f"{tag}-{i}",),
                    )
                    conn.commit()
                conn.close()
            except Exception as exc:  # noqa: BLE001
                errors.append(f"{tag}: {exc}")

        threads = [threading.Thread(target=writer, args=(f"PC{n}",)) for n in (1, 2, 3)]
        for t in threads:
            t.start()
        for t in threads:
            t.join()

        assert not errors
        conn = db.connect()
        assert conn.execute("SELECT COUNT(*) FROM requisitions").fetchone()[0] == 75
        conn.close()


class TestErrorPages:
    def test_missing_requisition_is_404(self, admin):
        assert admin.get("/requisition/9999").status_code == 404

    def test_missing_dish_is_404(self, admin):
        assert admin.get("/dishes/9999").status_code == 404

    def test_anonymous_user_is_sent_to_login(self, application):
        client = application.test_client()
        r = client.get("/")
        assert "/login" in r.headers["Location"]


class TestDocumentImportUI:
    """The in-app importer: an optional add-on that must stay optional.

    The point of these tests is the boundary. The importer may pre-fill the
    recipe form, but it must not acquire a second way into the database, and the
    app must keep working when the importer is not installed at all.
    """

    SAMPLE = os.path.join(
        os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
        "SampleIngredient.docx",
    )

    def test_app_starts_and_serves_without_the_importer(self, admin, monkeypatch):
        # Simulate a Mess PC where recipe_import/ was never copied across.
        import app as appmod

        monkeypatch.setattr(appmod, "_load_importer", lambda: None)
        monkeypatch.setattr(appmod, "importer_available", lambda: False)
        monkeypatch.setattr(
            appmod,
            "importer_error_message",
            lambda: "Recipe import is not installed on this PC.",
        )

        assert appmod.importer_available() is False
        assert admin.get("/dishes").status_code == 200
        # Button stays visible so admins can open the page and see install help.
        assert b"Import from document" in admin.get("/dishes").data
        blocked = admin.get("/dishes/import")
        assert blocked.status_code == 503
        assert b"not available" in blocked.data
        assert admin.get("/dishes/import/review").status_code == 404

    def test_import_page_is_offered_when_the_importer_is_present(self, admin):
        pytest.importorskip("recipe_import")
        assert b"Import from document" in admin.get("/dishes").data
        assert admin.get("/dishes/import").status_code == 200

    def test_staff_can_reach_the_importer(self, staff):
        pytest.importorskip("recipe_import")
        assert staff.get("/dishes/import").status_code == 200
        assert staff.post("/dishes/import/cancel", follow_redirects=True).status_code == 200

    def test_a_document_is_parsed_into_a_prefilled_form(self, admin):
        pytest.importorskip("recipe_import")
        if not os.path.exists(self.SAMPLE):
            pytest.skip("sample document absent")

        with open(self.SAMPLE, "rb") as fh:
            response = admin.post(
                "/dishes/import",
                data={"document": (io.BytesIO(fh.read()), "SampleIngredient.docx")},
                content_type="multipart/form-data",
                follow_redirects=True,
            )

        body = response.data.decode()
        assert response.status_code == 200
        assert "Recipe 1 of 2" in body
        assert "Paneer Kaleji" in body
        # The document's own wording travels with the row it produced.
        assert "30 g deggi mirch Salt to taste" in body
        # base_persons is never guessed - the field is left empty for a human.
        assert 'name="base_persons" min="1" required\n               value=""' in body \
            or 'value=""' in body

    def test_a_flagged_row_cannot_be_saved_through_the_import_form(self, admin):
        """The importer pre-fills the form; it does not bypass its validation."""
        pytest.importorskip("recipe_import")

        response = admin.post(
            "/dishes/new",
            data={
                "name": "Amritsari Paneer Bhurji",
                "category": "Main Course",
                "base_persons": "100",
                "import_next": "1",
                # "2-inch ginger" as it arrives: a quantity with no usable unit.
                "ingredient_name": ["Paneer", "Ginger"],
                "ingredient_quantity": ["1", "2"],
                "ingredient_unit": ["kg", ""],
            },
            follow_redirects=True,
        )

        assert b"is not a recognised unit" in response.data
        conn = db.connect()
        assert conn.execute("SELECT COUNT(*) FROM dishes").fetchone()[0] == 0
        conn.close()

    def test_import_validation_rerender_keeps_document_annotations(self, admin):
        """A bad field must not strip status/source_line from the import form."""
        pytest.importorskip("recipe_import")
        import json
        import app as appmod

        token = "testimporttoken1"
        payload = [
            {
                "dish_name": "Amritsari Paneer Bhurji",
                "base_persons": None,
                "category": "Main Course",
                "source_file": "sample.docx",
                "rows": [
                    {
                        "name": "Paneer",
                        "quantity": 1,
                        "unit": "kg",
                        "status": "ok",
                        "note": "",
                        "source_line": "1 kg Paneer",
                        "suggested_name": "",
                    },
                    {
                        "name": "Ginger",
                        "quantity": 2,
                        "unit": "",
                        "status": "unit_unconvertible",
                        "note": "length, not a store unit",
                        "source_line": "2-inch ginger",
                        "suggested_name": "",
                    },
                ],
            }
        ]
        path = appmod._import_store_path(token)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path, "w", encoding="utf-8") as fh:
            json.dump(payload, fh)

        with admin.session_transaction() as sess:
            sess["import_token"] = token
            sess["import_index"] = 0

        response = admin.post(
            "/dishes/new",
            data={
                "name": "Amritsari Paneer Bhurji",
                "category": "Main Course",
                "base_persons": "100",
                "import_next": "1",
                "ingredient_name": ["Paneer", "Ginger"],
                "ingredient_quantity": ["1", "2"],
                "ingredient_unit": ["kg", ""],
            },
            follow_redirects=True,
        )

        body = response.data.decode()
        assert "is not a recognised unit" in body
        # Flagged row and the document's own wording must still be on screen.
        assert "Ginger" in body
        assert "2-inch ginger" in body
        assert "length, not a store unit" in body
        assert "needs-attention" in body
        assert "could not be" in body  # flagged banner still counts rows

        try:
            os.remove(path)
        except OSError:
            pass

    def test_a_recipe_with_no_serving_count_is_refused(self, admin):
        pytest.importorskip("recipe_import")

        response = admin.post(
            "/dishes/new",
            data={
                "name": "Paneer Kaleji", "category": "Main Course",
                "base_persons": "", "import_next": "1",
                "ingredient_name": ["Paneer"], "ingredient_quantity": ["1"],
                "ingredient_unit": ["kg"],
            },
            follow_redirects=True,
        )

        assert b"whole number" in response.data
        conn = db.connect()
        assert conn.execute("SELECT COUNT(*) FROM dishes").fetchone()[0] == 0
        conn.close()

    def test_saving_moves_on_to_the_next_recipe_in_the_document(self, admin):
        pytest.importorskip("recipe_import")
        if not os.path.exists(self.SAMPLE):
            pytest.skip("sample document absent")

        with open(self.SAMPLE, "rb") as fh:
            admin.post(
                "/dishes/import",
                data={"document": (io.BytesIO(fh.read()), "SampleIngredient.docx")},
                content_type="multipart/form-data",
            )

        saved = admin.post(
            "/dishes/new",
            data={
                "name": "Paneer Kaleji", "category": "Main Course",
                "base_persons": "100", "import_next": "1",
                "ingredient_name": ["Paneer"], "ingredient_quantity": ["1"],
                "ingredient_unit": ["kg"],
            },
        )
        assert saved.headers["Location"].endswith("/dishes/import/review")

        # ...and the next screen is the document's second recipe, not the first.
        body = admin.get("/dishes/import/review").data.decode()
        assert "Recipe 2 of 2" in body
        assert "Amritsari Paneer Bhurji" in body

    def test_abandoning_an_import_saves_nothing(self, admin):
        pytest.importorskip("recipe_import")
        if not os.path.exists(self.SAMPLE):
            pytest.skip("sample document absent")

        with open(self.SAMPLE, "rb") as fh:
            admin.post(
                "/dishes/import",
                data={"document": (io.BytesIO(fh.read()), "SampleIngredient.docx")},
                content_type="multipart/form-data",
            )

        admin.post("/dishes/import/cancel", follow_redirects=True)

        conn = db.connect()
        assert conn.execute("SELECT COUNT(*) FROM dishes").fetchone()[0] == 0
        conn.close()
        assert admin.get("/dishes/import/review", follow_redirects=True).status_code == 200

    def test_a_rejected_file_type_does_not_reach_the_parser(self, admin):
        pytest.importorskip("recipe_import")
        response = admin.post(
            "/dishes/import",
            data={"document": (io.BytesIO(b"rm -rf /"), "payload.exe")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        assert b"not a document type" in response.data

    def test_an_unreadable_document_does_not_crash_the_app(self, admin):
        pytest.importorskip("recipe_import")
        response = admin.post(
            "/dishes/import",
            data={"document": (io.BytesIO(b"not really a docx"), "broken.docx")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
        assert response.status_code == 200
        assert b"Could not read that document" in response.data


DAL = {
    "name": "Dal Tadka",
    "category": "Dal & Vegetables",
    "base_persons": "10",
    "notes": "",
    "ingredient_name": ["Toor Dal", "onion", "Salt"],
    "ingredient_quantity": ["1", "250", "50"],
    "ingredient_unit": ["kg", "g", "g"],
}


def _meal(client, dish_ids, persons=None, default="90", **extra):
    persons = persons or [""] * len(dish_ids)
    data = {"meal_type": "Lunch", "meal_date": "2026-09-26", "course_name": "FC 2026",
            "default_persons": default, "dish_id": dish_ids, "dish_persons": persons}
    data.update(extra)
    return client.post("/meal/new", data=data, follow_redirects=True)


def _count(table):
    conn = db.connect()
    n = conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
    conn.close()
    return n


class TestMealIndent:
    @pytest.fixture(autouse=True)
    def _recipes(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        admin.post("/dishes/new", data=DAL)
        self.admin = admin

    def test_one_meal_writes_a_requisition_per_dish(self):
        r = _meal(self.admin, ["1", "2"])
        assert r.status_code == 200
        assert _count("meal_indents") == 1
        conn = db.connect()
        rows = conn.execute(
            "SELECT dish_name_snapshot, persons_required, meal_indent_id FROM requisitions"
        ).fetchall()
        conn.close()
        assert {tuple(row) for row in rows} == {
            ("Chicken Biryani", 90, 1), ("Dal Tadka", 90, 1)}

    def test_shared_ingredients_are_totalled(self):
        html = _meal(self.admin, ["1", "2"]).data.decode()
        # Onion: 13.5 kg (biryani) + 2.25 kg (dal, 2250 g) = 15.75 kg
        assert "15.75" in html
        # Salt: 900 g + 450 g = 1.35 kg
        assert "1.35" in html
        assert "Chicken Biryani 13.5 kg; Dal Tadka 2.25 kg" in html

    def test_per_dish_override_beats_the_meal_strength(self):
        html = _meal(self.admin, ["1", "2"], persons=["", "30"]).data.decode()
        # Onion: 13.5 kg + 750 g = 14.25 kg
        assert "14.25" in html
        conn = db.connect()
        dal = conn.execute(
            "SELECT persons_required FROM requisitions WHERE dish_name_snapshot = 'Dal Tadka'"
        ).fetchone()
        conn.close()
        assert dal[0] == 30

    def test_blank_rows_are_ignored(self):
        _meal(self.admin, ["1", "", "2", ""])
        assert _count("requisitions") == 2

    def test_duplicate_dish_is_rejected_and_nothing_written(self):
        r = _meal(self.admin, ["1", "1"])
        assert b"selected more than once" in r.data
        assert _count("meal_indents") == 0 and _count("requisitions") == 0

    def test_soft_deleted_dish_is_rejected(self):
        self.admin.post("/dishes/2/delete")
        r = _meal(self.admin, ["1", "2"])
        assert b"valid dish" in r.data
        assert _count("requisitions") == 0

    def test_at_least_one_dish_required(self):
        r = _meal(self.admin, ["", ""])
        assert b"at least one dish" in r.data

    def test_bad_strength_and_override_rejected(self):
        assert b"at least 1" in _meal(self.admin, ["1"], default="0").data
        assert b"whole number" in _meal(self.admin, ["1"], persons=["abc"]).data
        assert _count("meal_indents") == 0

    def test_input_survives_a_validation_failure(self):
        html = _meal(self.admin, ["1", "1"], persons=["", "45"], default="77").data.decode()
        assert 'value="77"' in html and 'value="45"' in html

    def test_failure_mid_meal_leaves_nothing_behind(self, monkeypatch):
        real_execute = db.execute
        calls = {"n": 0}

        def flaky_execute(sql, args=()):
            if "INSERT INTO requisitions" in sql:
                calls["n"] += 1
                if calls["n"] == 2:
                    raise sqlite3.OperationalError("simulated failure on the second dish")
            return real_execute(sql, args)

        monkeypatch.setattr(db, "execute", flaky_execute)
        with pytest.raises(sqlite3.OperationalError):
            _meal(self.admin, ["1", "2"])
        assert _count("meal_indents") == 0
        assert _count("requisitions") == 0
        assert _count("requisition_items") == 0

    def test_editing_a_recipe_does_not_change_an_issued_meal(self):
        _meal(self.admin, ["1", "2"])
        amended = dict(DAL, ingredient_quantity=["1", "500", "50"])
        self.admin.post("/dishes/2/edit", data=amended)
        assert "15.75" in self.admin.get("/meal/1").data.decode()

    def test_meal_pdf_holds_sheet_and_slips(self):
        _meal(self.admin, ["1", "2"])
        r = self.admin.get("/meal/1/pdf")
        assert r.status_code == 200 and r.data[:4] == b"%PDF"
        assert r.data.count(b"/Type /Page\n") >= 3 or b"Page 3" in r.data

    def test_meal_excel_has_consolidated_and_dish_sheets(self):
        from openpyxl import load_workbook

        _meal(self.admin, ["1", "2"])
        wb = load_workbook(io.BytesIO(self.admin.get("/meal/1/excel").data))
        assert wb.sheetnames == ["Consolidated", "Chicken Biryani", "Dal Tadka"]
        ws = wb["Consolidated"]
        onion = next(
            [ws.cell(row=rw, column=c).value for c in range(1, 7)]
            for rw in range(1, ws.max_row + 1)
            if str(ws.cell(row=rw, column=2).value).casefold() == "onion"
        )
        assert onion[2] == 15.75 and onion[3] == "kg"

    def test_missing_meal_is_404(self):
        assert self.admin.get("/meal/99").status_code == 404

    def test_history_lists_the_meal(self):
        _meal(self.admin, ["1", "2"])
        html = self.admin.get("/history").data.decode()
        assert "Meal indents" in html and "Meal #1" in html

    def test_staff_can_build_a_meal(self, staff):
        r = _meal(staff, ["1", "2"])
        assert r.status_code == 200 and b"15.75" in r.data

    def test_single_dish_calculate_is_not_part_of_a_meal(self):
        self.admin.post("/calculate", data={"dish_id": "1", "persons_required": "90",
                                            "meal_type": "Lunch", "meal_date": "2026-08-10"})
        conn = db.connect()
        assert conn.execute("SELECT meal_indent_id FROM requisitions").fetchone()[0] is None
        conn.close()


class TestSavedMenus:
    @pytest.fixture(autouse=True)
    def _recipes(self, admin):
        admin.post("/dishes/new", data=BIRYANI)
        admin.post("/dishes/new", data=DAL)
        self.admin = admin

    def test_save_and_load(self):
        r = _meal(self.admin, ["1", "2"], persons=["", "30"], save_menu_name="Sunday Lunch")
        assert b"Sunday Lunch&#39; saved" in r.data
        html = self.admin.get("/meal/new?menu_id=1").data.decode()
        assert 'value="90"' in html and 'value="30"' in html
        assert "Sunday Lunch" in self.admin.get("/menus").data.decode()

    def test_saving_under_the_same_name_replaces(self):
        _meal(self.admin, ["1", "2"], save_menu_name="Sunday Lunch")
        _meal(self.admin, ["2"], save_menu_name="sunday lunch")
        assert _count("menus") == 1
        assert _count("menu_dishes") == 1

    def test_soft_deleted_dish_is_skipped_on_load(self):
        _meal(self.admin, ["1", "2"], save_menu_name="Sunday Lunch")
        self.admin.post("/dishes/2/delete")
        html = self.admin.get("/meal/new?menu_id=1").data.decode()
        assert "left out of this menu" in html

    def test_delete_menu_keeps_issued_meals(self):
        _meal(self.admin, ["1", "2"], save_menu_name="Sunday Lunch")
        self.admin.post("/menus/1/delete")
        assert _count("menus") == 0 and _count("menu_dishes") == 0
        assert self.admin.get("/meal/1").status_code == 200

    def test_failed_meal_does_not_save_the_menu(self):
        _meal(self.admin, ["1", "1"], save_menu_name="Broken")
        assert _count("menus") == 0


class TestMigration:
    def test_old_database_gains_meal_indent_column(self, tmp_path, monkeypatch):
        path = tmp_path / "old.db"
        conn = sqlite3.connect(path)
        # The requisitions table as it was before meal indents existed.
        conn.executescript("""
            CREATE TABLE requisitions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                dish_id INTEGER, dish_name_snapshot TEXT NOT NULL,
                base_persons_snapshot INTEGER NOT NULL,
                persons_required INTEGER NOT NULL, course_name TEXT NOT NULL DEFAULT '',
                meal_type TEXT NOT NULL, meal_date TEXT NOT NULL,
                generated_by INTEGER, generated_by_name TEXT NOT NULL DEFAULT '',
                generated_at TEXT NOT NULL DEFAULT '');
            INSERT INTO requisitions (dish_name_snapshot, base_persons_snapshot,
                persons_required, meal_type, meal_date)
            VALUES ('Old Dish', 10, 90, 'Lunch', '2026-01-01');
        """)
        conn.commit()
        conn.close()

        monkeypatch.setattr(db, "DB_PATH", str(path))
        db.init_db()
        db.init_db()  # safe to re-run

        conn = db.connect()
        columns = {r["name"] for r in conn.execute("PRAGMA table_info(requisitions)")}
        old = conn.execute(
            "SELECT dish_name_snapshot, meal_indent_id FROM requisitions").fetchone()
        conn.close()
        assert "meal_indent_id" in columns
        assert tuple(old) == ("Old Dish", None)
